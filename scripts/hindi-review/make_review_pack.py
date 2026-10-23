# Build the Hindi review workbook for a native reader: each top page's English beside its Hindi (from the built
# site), text shared by many pages once on its own sheet, the September capsule titles, and the SSC glossary
# (src/data/hindiGlossary.mjs).
#
# Usage, after `npm run build` (with the server's current data):
#   python3 scripts/hindi-review/make_review_pack.py <server prisma/seed-data/gpt-current-affairs dir> <out.xlsx>
# Needs openpyxl (pip install openpyxl) and node.
import collections, difflib, glob, json, os, re, subprocess, sys
sys.path.insert(0, os.path.dirname(__file__))
from blocks import blocks
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), '..', '..'))
DIST = os.path.join(ROOT, 'dist')
SEED = sys.argv[1]
CATS = os.path.join(ROOT, 'scripts', 'capsule-pdf', 'categories.json')
OUT = sys.argv[2]
SITE = 'https://parikshasaathi.com'
DEVA = re.compile('[ऀ-ॿ]')
MONTHS = re.compile(r'\b(?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)[a-z]*\b')

HUBS = ['ssc-cgl', 'ssc-chsl', 'ssc-mts', 'ssc-gd', 'ibps-po', 'ibps-clerk', 'sbi-po', 'sbi-clerk', 'rrb-ntpc', 'rrb-group-d',
        'rrb-alp', 'agniveer-army', 'ctet', 'up-police', 'bihar-police', 'mp-police', 'rajasthan-police', 'bpsc', 'mppsc', 'rpsc', 'uppsc']
PAGES = [('Hindi home page', '/', '/hi/'), ('Exam chooser', '/exams/', '/hi/exams/'), ('Exam calendar', '/exams/calendar/', '/hi/exams/calendar/')]
PAGES += [(None, f'/exams/{h}/', f'/hi/exams/{h}/') for h in HUBS]

TAG = {'h1': 'Page heading', 'h2': 'Heading', 'h3': 'Sub-heading', 'h4': 'Sub-heading', 'p': 'Paragraph', 'li': 'List item',
       'td': 'Table cell', 'th': 'Table heading', 'dt': 'Label', 'dd': 'Value', 'summary': 'Heading', 'caption': 'Table title',
       'figcaption': 'Caption', 'blockquote': 'Quote'}

def pair(en, hi):
    if len(en) == len(hi):
        return list(zip(en, hi))
    out = []
    sm = difflib.SequenceMatcher(None, [t for t, _ in en], [t for t, _ in hi], autojunk=False)
    for op, i1, i2, j1, j2 in sm.get_opcodes():
        a, b = en[i1:i2], hi[j1:j2]
        for k in range(max(len(a), len(b))):
            out.append((a[k] if k < len(a) else (b[k][0], ''), b[k] if k < len(b) else (a[k][0], '')))
    return out

page_rows = []
for name, en_path, hi_path in PAGES:
    en, hi = blocks(f'{DIST}{en_path}index.html'), blocks(f'{DIST}{hi_path}index.html')
    if name is None:
        name = en[0][1].split(':')[0].replace(' 2026', '').replace(' 2025', '').strip()
    rows, seen, templates = [], set(), {}
    for (tag, en_text), (_, hi_text) in pair(en, hi):
        if not DEVA.search(hi_text) or (en_text, hi_text) in seen:
            continue
        seen.add((en_text, hi_text))
        # Rows that differ only in numbers and month names (the solved-paper lists) are reviewed once.
        template = (tag, MONTHS.sub('M', re.sub(r'[0-9]+', '#', en_text)))
        if template in templates:
            templates[template][1] += 1
            continue
        templates[template] = [len(rows), 0]
        rows.append([TAG.get(tag, tag), en_text, hi_text])
    for index, extra in templates.values():
        if extra:
            rows[index][0] += f' · +{extra} similar'
    page_rows.append((name, SITE + hi_path, [tuple(r) for r in rows]))

# Text on three or more pages is reviewed once, on the shared sheet.
counts = collections.Counter((r[1], r[2]) for _, _, rows in page_rows for r in rows)
shared_keys = {k for k, n in counts.items() if n >= 3}
shared = {}
for name, _, rows in page_rows:
    for section, en_text, hi_text in rows:
        if (en_text, hi_text) in shared_keys:
            shared.setdefault((en_text, hi_text), [section, set()])[1].add(name)

# September 2026 capsule: one title per item, paired by the official release it came from.
cats = json.load(open(CATS, encoding='utf-8'))
themes = [('economy', 'Economy, Banking & Finance', 'अर्थव्यवस्था, बैंकिंग व वित्त'), ('governance', 'Schemes, Governance & Appointments', 'योजनाएँ, शासन व नियुक्तियाँ'),
          ('defence', 'Defence & Security', 'रक्षा व सुरक्षा'), ('international', 'International Relations', 'अंतरराष्ट्रीय संबंध'),
          ('science', 'Science & Technology', 'विज्ञान व प्रौद्योगिकी'), ('agri-env', 'Environment & Agriculture', 'पर्यावरण व कृषि'),
          ('health', 'Health', 'स्वास्थ्य'), ('culture', 'Sports, Awards & Culture', 'खेल, पुरस्कार व संस्कृति')]
items = collections.defaultdict(dict)
for path in glob.glob(os.path.join(SEED, '*.json')):
    for it in json.load(open(path, encoding='utf-8')).get('currentAffairs', []):
        d = it.get('editionDate', '')[:10]
        if '2026-09-12' <= d <= '2026-09-30':
            items[it['sourceUrl']][it.get('languageCode', 'en')] = it
capsule = []
for url, by_lang in items.items():
    m = re.search(r'PRID=(\d+)', url)
    theme = cats.get(m.group(1)) if m else None
    if theme and 'en' in by_lang and 'hi' in by_lang:
        capsule.append((theme, by_lang['en']['editionDate'][:10], by_lang['en']['title'], by_lang['hi']['title']))
order = {k: i for i, (k, _, _) in enumerate(themes)}
capsule.sort(key=lambda r: (order[r[0]], r[1]))
theme_names = {k: (en, hi) for k, en, hi in themes}

gloss = json.loads(subprocess.run(
    ['node', '--input-type=module', '-e',
     "import * as g from './src/data/hindiGlossary.mjs'; console.log(JSON.stringify({ sources: g.HINDI_GLOSSARY_SOURCES, terms: g.HINDI_GLOSSARY, subjects: g.SSC_SUBJECTS_HI, posts: g.SSC_POSTS_HI, readOn: g.HINDI_GLOSSARY_READ_ON }))"],
    cwd=ROOT, check=True, capture_output=True, text=True).stdout)

# ---- workbook ----
wb = Workbook()
EN_FONT, HI_FONT = Font(name='Arial', size=10), Font(name='Nirmala UI', size=11)
HEAD = Font(name='Arial', size=10, bold=True, color='FFFFFF')
HEAD_FILL = PatternFill('solid', fgColor='1F4E5F')
READER_FILL = PatternFill('solid', fgColor='FFF7D6')
TITLE = Font(name='Arial', size=14, bold=True, color='1F4E5F')
SUB = Font(name='Arial', size=10, italic=True, color='555555')
THIN = Border(bottom=Side(style='thin', color='D0D7DA'))
WRAP = Alignment(wrap_text=True, vertical='top')
VERDICTS = '"OK,Fix,Unsure"'

def header(ws, row, labels, widths):
    for c, (label, width) in enumerate(zip(labels, widths), 1):
        cell = ws.cell(row=row, column=c, value=label)
        cell.font, cell.fill, cell.alignment = HEAD, HEAD_FILL, Alignment(wrap_text=True, vertical='center')
        ws.column_dimensions[cell.column_letter].width = width
    ws.freeze_panes = ws.cell(row=row + 1, column=1)

def reader_columns(ws, first_row, last_row, verdict_col):
    dv = DataValidation(type='list', formula1=VERDICTS, allow_blank=True)
    ws.add_data_validation(dv)
    if last_row >= first_row:
        dv.add(f'{ws.cell(row=first_row, column=verdict_col).coordinate}:{ws.cell(row=last_row, column=verdict_col).coordinate}')
    for r in range(first_row, last_row + 1):
        for c in (verdict_col, verdict_col + 1, verdict_col + 2):
            cell = ws.cell(row=r, column=c)
            cell.fill = READER_FILL
            cell.font = HI_FONT if c == verdict_col + 1 else EN_FONT
            cell.alignment = WRAP
            cell.border = THIN

def write_rows(ws, start, rows, hi_cols):
    for i, values in enumerate(rows):
        for c, v in enumerate(values, 1):
            cell = ws.cell(row=start + i, column=c, value=v)
            cell.font = HI_FONT if c in hi_cols else EN_FONT
            cell.alignment = WRAP
            cell.border = THIN

# Start here
ws = wb.active
ws.title = 'Start here'
ws.column_dimensions['A'].width = 30
ws.column_dimensions['B'].width = 95
ws['A1'] = 'Pariksha Saathi: Hindi review pack'
ws['A1'].font = TITLE
ws['A2'] = f'Built on {gloss["readOn"]} from the site\'s text as it reads after that day\'s Hindi glossary update. Please return this file with the yellow columns filled in.'
ws['A2'].font = SUB
lines = [
    ('What to check', 'Is each Hindi line correct, natural and clear for a Hindi-medium exam aspirant, and does it say the same as the English beside it?'),
    ('How to review', 'Open each page sheet in order. For every row choose a verdict in the yellow "Verdict" column. If you choose Fix, write your Hindi in "Your Hindi". Use "Note" for anything else.'),
    ('Verdicts', 'OK = correct as it is. Fix = needs a change (write it). Unsure = flag for discussion.'),
    ('Shared text', 'Text that appears on three or more pages is on the "Shared text" sheet only, so you review it once.'),
    ('Official terms', 'Exam terms follow SSC\'s own Hindi notices (see the "Glossary" sheet): अभ्यर्थी, प्रवेश पत्र, उत्तर कुंजी, नकारात्मक अंकन, पाठ्यक्रम, वेतन स्तर, भूतपूर्व सैनिक, कौशल परीक्षा. These are deliberate; flag one only if you think SSC itself would not use it.'),
    ('Search words', 'Page titles and some headings keep the words students type into Google (एडमिट कार्ड, सिलेबस, नेगेटिव मार्किंग), sometimes in brackets after the official term. That is deliberate too.'),
    ('Names and numbers', 'Exam names, post names, dates and figures were checked against the official documents. Please do not change numbers; flag them in "Note" if one looks wrong.'),
    ('Example row', 'English: "Negative marking"  ·  Hindi: "नकारात्मक अंकन"  ·  Verdict: OK. Or: Hindi "अंतिम सत्यापन"  ·  Verdict: Fix  ·  Your Hindi: "अंतिम बार जाँचा गया"  ·  Note: "सत्यापन sounds formal here".'),
]
for i, (k, v) in enumerate(lines, 4):
    ws.cell(row=i, column=1, value=k).font = Font(name='Arial', size=10, bold=True)
    c = ws.cell(row=i, column=2, value=v)
    c.font, c.alignment = (HI_FONT if DEVA.search(v) else EN_FONT), WRAP
    ws.cell(row=i, column=1).alignment = WRAP
r = 4 + len(lines) + 1
ws.cell(row=r, column=1, value='Sheets').font = Font(name='Arial', size=11, bold=True)
contents_row = r + 1

# Glossary
g = wb.create_sheet('Glossary')
g['A1'] = 'Exam terms as SSC prints them in its Hindi notices'
g['A1'].font = TITLE
g['A2'] = 'Read from the notices\' page images on ' + gloss['readOn'] + '. "Avoid" forms stay only in titles, link text, or in brackets after the official term.'
g['A2'].font = SUB
header(g, 4, ['English', 'Official Hindi (used on the site)', 'Also printed by SSC', 'Avoid in body text', 'Source notice and pages'],
       [30, 34, 40, 28, 52])
src = gloss['sources']
def src_text(pairs):
    return '; '.join(f'{src[k]["en"].split(":")[0]} p. {", ".join(map(str, pages))}' for k, pages in pairs)
rows = [(t['en'], t['hi'], '; '.join(t.get('also', [])), ', '.join(a.replace('(?<!भूत)', '') for a in t.get('avoid', [])), src_text(t['src']))
        for t in gloss['terms']]
write_rows(g, 5, rows, {2, 3, 4})
r = 5 + len(rows) + 1
g.cell(row=r, column=1, value='Subject names, exam by exam (each exam\'s own notice)').font = Font(name='Arial', size=11, bold=True)
r += 1
exam_names = {'ssc_cgl': 'SSC CGL', 'ssc_chsl': 'SSC CHSL', 'ssc_mts': 'SSC MTS', 'ssc_gd': 'SSC GD'}
for exam, s in gloss['subjects'].items():
    for en_name, hi_name in s['names'].items():
        write_rows(g, r, [(f'{exam_names[exam]}: {en_name}', hi_name, '', '', f'{src[s["src"]]["en"].split(":")[0]} p. {", ".join(map(str, s["pages"]))}')], {2})
        r += 1
    for en_name, hi_name in s.get('unofficial', {}).items():
        write_rows(g, r, [(f'{exam_names[exam]}: {en_name}', hi_name, '', '', 'Not in any SSC Hindi notice yet (new in CGL 2026): our translation, please check')], {2})
        r += 1
r += 1
g.cell(row=r, column=1, value='Post titles (CGL 2025 and CHSL 2026 notices)').font = Font(name='Arial', size=11, bold=True)
r += 1
for en_name, hi_name in gloss['posts'].items():
    write_rows(g, r, [(en_name, hi_name, '', '', '')], {2})
    r += 1
r += 1
g.cell(row=r, column=1, value='Sources').font = Font(name='Arial', size=11, bold=True)
r += 1
for k, s in src.items():
    write_rows(g, r, [(s['en'], s['hi'], f'dated {s["dated"]}', '', s['url'])], {2})
    r += 1

COLS = ['#', 'Where', 'English', 'Hindi on the site', 'Verdict', 'Your Hindi', 'Note']
WIDTHS = [5, 14, 52, 58, 10, 46, 30]
contents = []

# Shared text
sh = wb.create_sheet('Shared text')
sh['A1'] = 'Text shared by three or more pages (review once)'
sh['A1'].font = TITLE
sh['A2'] = 'The "Where" column says how many of the pages carry the line.'
sh['A2'].font = SUB
header(sh, 4, COLS, WIDTHS)
shared_rows = sorted(shared.items(), key=lambda kv: (-len(kv[1][1]), kv[0][0]))
write_rows(sh, 5, [(i, f'{v[0]} · {len(v[1])} pages', k[0], k[1]) for i, (k, v) in enumerate(shared_rows, 1)], {4})
reader_columns(sh, 5, 4 + len(shared_rows), 5)
contents.append(('Shared text', len(shared_rows), ''))

# Page sheets
for n, (name, url, rows) in enumerate(page_rows, 1):
    rows = [r for r in rows if (r[1], r[2]) not in shared_keys]
    title = f'{n:02d} {name}'[:31].replace('/', '-')
    ws_p = wb.create_sheet(title)
    ws_p['A1'] = name
    ws_p['A1'].font = TITLE
    ws_p['A2'] = url
    ws_p['A2'].font = SUB
    ws_p['A2'].hyperlink = url
    header(ws_p, 4, COLS, WIDTHS)
    write_rows(ws_p, 5, [(i, sec, en_t, hi_t) for i, (sec, en_t, hi_t) in enumerate(rows, 1)], {4})
    reader_columns(ws_p, 5, 4 + len(rows), 5)
    contents.append((title, len(rows), url))

# Capsule
cp = wb.create_sheet(f'{len(page_rows) + 1:02d} Sept capsule PDF')
cp['A1'] = 'Current Affairs Capsule, September 2026 (Hindi PDF)'
cp['A1'].font = TITLE
cp['A2'] = SITE + '/pdf/ParikshaSaathi_Capsule_2026-09_hi.pdf'
cp['A2'].font = SUB
cp['A2'].hyperlink = SITE + '/pdf/ParikshaSaathi_Capsule_2026-09_hi.pdf'
header(cp, 4, ['#', 'Theme', 'English title', 'Hindi title in the PDF', 'Verdict', 'Your Hindi', 'Note'], WIDTHS)
cap_rows = [(i, f'{theme_names[t][0]} · {d[8:10]}/{d[5:7]}', en_t, hi_t) for i, (t, d, en_t, hi_t) in enumerate(capsule, 1)]
cap_rows += [(len(cap_rows) + i, 'Theme heading', en, hi) for i, (_, en, hi) in enumerate(themes, 1)]
write_rows(cp, 5, cap_rows, {4})
reader_columns(cp, 5, 4 + len(cap_rows), 5)
contents.append((cp.title, len(cap_rows), cp['A2'].value))

# Contents on the start sheet
ws.cell(row=contents_row, column=1, value='Sheet').font = Font(name='Arial', size=10, bold=True)
ws.cell(row=contents_row, column=2, value='Rows to review · page').font = Font(name='Arial', size=10, bold=True)
total = 0
for i, (title, n, url) in enumerate(contents, contents_row + 1):
    c = ws.cell(row=i, column=1, value=title)
    c.font = Font(name='Arial', size=10, color='1F4E5F', underline='single')
    c.hyperlink = f"#'{title}'!A1"
    ws.cell(row=i, column=2, value=f'{n} rows' + (f' · {url}' if url else '')).font = EN_FONT
    total += n
ws.cell(row=contents_row + len(contents) + 1, column=1, value='Total').font = Font(name='Arial', size=10, bold=True)
ws.cell(row=contents_row + len(contents) + 1, column=2, value=f'{total} rows (counted when this file was built)').font = EN_FONT

os.makedirs(os.path.dirname(OUT), exist_ok=True)
wb.save(OUT)
print('saved', OUT, '| sheets', len(wb.sheetnames), '| rows', total, '| shared', len(shared_rows), '| capsule items', len(capsule))
for title, n, _ in contents: print(f'  {title}: {n}')
